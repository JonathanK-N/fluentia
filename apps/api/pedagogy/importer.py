"""Modèles Excel des quatre épreuves et import d’une série complète (tableur seul ou .zip avec médias).

Un fichier = une série = une épreuve. Rien n’est enregistré tant qu’une seule ligne est en erreur.
"""

import io
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from rest_framework.exceptions import ValidationError

from exams.models import LEVEL_POINTS
from .models import SeriesImport, TaskSet, TaskSubject
from .services import (
    check_audio, check_image, create_item, event, image_path, media_path, require, store_audio, store_image,
)

SKILLS = {"reading": "Lecture", "listening": "Écoute", "writing": "Rédaction", "speaking": "Oral"}
LABELS = {label.casefold(): skill for skill, label in SKILLS.items()}
LEVELS = list(LEVEL_POINTS)
LETTERS = "ABCD"
QCM = ("reading", "listening")
DURATIONS = {"reading": 60, "listening": 35}  # minutes, format de simulation
MAX_ROWS = 100
MAX_ARCHIVE = 250_000_000

CHOICES = ["Choix A", "Choix B", "Choix C", "Choix D", "Bonne réponse", "Indice"]
COLUMNS = {
    "reading": ["N°", "Niveau", "Texte", "Image", "Texte de l’image", "Question", *CHOICES],
    "listening": ["N°", "Niveau", "Audio", "Transcription", "Image", "Texte de l’image", "Question", *CHOICES],
    "writing": ["Tâche", "Titre du sujet", "Consigne", "Type de texte", "Mots minimum", "Mots maximum", "Durée (minutes)", "Document 1", "Document 2"],
    "speaking": ["Tâche", "Titre du sujet", "Consigne", "Préparation (minutes)", "Temps de parole (minutes)", "Pistes"],
}
WIDTHS = {"N°": 6, "Tâche": 8, "Niveau": 9, "Texte": 60, "Transcription": 60, "Consigne": 60, "Question": 40, "Indice": 34,
          "Texte de l’image": 40, "Document 1": 50, "Document 2": 50, "Pistes": 40, "Bonne réponse": 14}

SERIES_ROWS = {
    "reading": ["Épreuve", "Titre", "Mode", "Description", "Provenance et droits", "Auteur ou autrice"],
    "listening": ["Épreuve", "Titre", "Mode", "Description", "Provenance et droits", "Auteur ou autrice"],
    "writing": ["Épreuve", "Titre", "Description", "Provenance et droits", "Auteur ou autrice"],
    "speaking": ["Épreuve", "Titre", "Description", "Provenance et droits", "Auteur ou autrice"],
}

GUIDE = {
    "reading": [
        "Une ligne par question, dans l’ordre (N° 1, 2, 3…). 39 questions pour une simulation.",
        "Texte : le document à lire. Image : le nom d’un fichier PNG ou JPEG (ex. q12-lettre.png), si le document est une capture.",
        "Il faut un Texte, une Image, ou les deux. Avec une image, « Texte de l’image » est obligatoire (lecteurs d’écran).",
        "Bonne réponse : A, B, C ou D. Indice : l’extrait du texte qui justifie la réponse.",
        "Avec des images : mettez le tableur et les images dans un même dossier compressé (.zip), puis importez le .zip.",
    ],
    "listening": [
        "Une ligne par question, dans l’ordre (N° 1, 2, 3…). 39 questions pour une simulation.",
        "Audio : le nom du fichier WAV du document (ex. q01-annonce.wav). Transcription : le texte entendu, obligatoire.",
        "Image : facultative (ex. une affiche liée au document). Avec une image, « Texte de l’image » est obligatoire.",
        "Bonne réponse : A, B, C ou D. Indice : le passage de l’audio qui justifie la réponse.",
        "Mettez le tableur, les audios et les images dans un même dossier compressé (.zip), puis importez le .zip.",
    ],
    "writing": [
        "Une ligne par sujet. Les 3 tâches doivent avoir au moins un sujet chacune.",
        "Type de texte, mots minimum et maximum, durée : identiques pour tous les sujets d’une même tâche.",
        "Document 1 et Document 2 : seulement pour la tâche 3 (les deux points de vue à comparer).",
        "Aucune bonne réponse : seuls les sujets sont importés.",
    ],
    "speaking": [
        "Une ligne par sujet. Les 3 tâches doivent avoir au moins un sujet chacune.",
        "Préparation et temps de parole en minutes (ex. 2 ou 5,5) : identiques pour tous les sujets d’une même tâche.",
        "Pistes : une piste par ligne dans la cellule (Alt + Entrée), ou séparées par des points-virgules.",
        "Aucune bonne réponse : seuls les sujets sont importés.",
    ],
}

INDIGO = "3B33A6"
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor=INDIGO)
SOFT = PatternFill("solid", fgColor="EEEBFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def build_template(skill):
    """Modèle Excel prêt à remplir : Série, Contenu (en-têtes et listes) et Mode d’emploi."""
    book = Workbook()
    series = book.active
    series.title = "Série"
    series.append(["Champ", "Valeur"])
    for label in SERIES_ROWS[skill]:
        series.append([label, SKILLS[skill] if label == "Épreuve" else ("Simulation" if label == "Mode" else "")])
    for cell in series[1]:
        cell.font, cell.fill = HEAD, FILL
    for row in series.iter_rows(min_row=2, max_col=1):
        row[0].font, row[0].fill = Font(bold=True), SOFT
    series.column_dimensions["A"].width, series.column_dimensions["B"].width = 24, 70
    epreuve = DataValidation(type="list", formula1='"Lecture,Écoute,Rédaction,Oral"', allow_blank=False)
    series.add_data_validation(epreuve)
    epreuve.add("B2")
    if skill in QCM:
        mode = DataValidation(type="list", formula1='"Simulation,Entraînement"', allow_blank=False)
        series.add_data_validation(mode)
        mode.add("B4")

    content = book.create_sheet("Contenu")
    columns = COLUMNS[skill]
    content.append(columns)
    for index, name in enumerate(columns, 1):
        cell = content.cell(row=1, column=index)
        cell.font, cell.fill, cell.alignment = HEAD, FILL, Alignment(wrap_text=True, vertical="center")
        content.column_dimensions[get_column_letter(index)].width = WIDTHS.get(name, 22)
    content.freeze_panes = "A2"
    if skill in QCM:
        level = DataValidation(type="list", formula1='"' + ",".join(LEVELS) + '"')
        answer = DataValidation(type="list", formula1='"A,B,C,D"')
        content.add_data_validation(level)
        content.add_data_validation(answer)
        level.add(f"B2:B{MAX_ROWS + 1}")
        answer.add(f"{get_column_letter(columns.index('Bonne réponse') + 1)}2:{get_column_letter(columns.index('Bonne réponse') + 1)}{MAX_ROWS + 1}")
    else:
        task = DataValidation(type="list", formula1='"1,2,3"')
        content.add_data_validation(task)
        task.add(f"A2:A{MAX_ROWS + 1}")

    guide = book.create_sheet("Mode d’emploi")
    guide.column_dimensions["A"].width = 110
    guide.append([f"Modèle {SKILLS[skill]} — Fluentia"])
    guide["A1"].font = Font(bold=True, size=14, color=INDIGO)
    guide.append(["1. Remplissez l’onglet « Série » (titre, provenance et droits des documents)."])
    guide.append(["2. Remplissez l’onglet « Contenu », sans modifier la ligne des titres de colonnes."])
    for line in GUIDE[skill]:
        guide.append(["• " + line])
    guide.append(["3. Dans Fluentia, ouvrez Édition › Importer une série, choisissez l’épreuve et déposez le fichier."])
    if skill in QCM:
        guide.append([""])
        guide.append(["Barème : les points d’une question dépendent de son niveau."])
        guide[f"A{guide.max_row}"].font = Font(bold=True)
        for level, points in LEVEL_POINTS.items():
            guide.append([f"{level} : {points} points"])
    for row in guide.iter_rows():
        row[0].alignment = WRAP
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


class ImportRejected(Exception):
    """Fichier refusé : la liste des erreurs est renvoyée telle quelle (lignes numériques)."""

    def __init__(self, errors):
        super().__init__(f"{len(errors)} erreur(s)")
        self.payload = {"detail": f"{len(errors)} erreur(s) : rien n’a été importé.", "errors": errors}


@dataclass
class Report:
    errors: list = field(default_factory=list)

    def add(self, message, line=None, column=None):
        if len(self.errors) < 200:
            self.errors.append({"ligne": line, "colonne": column, "message": message})

    def raise_if_any(self):
        if self.errors:
            raise ImportRejected(self.errors)


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def number(value):
    raw = text(value).replace(",", ".")
    try:
        return float(raw)
    except ValueError:
        return None


def open_upload(uploaded, report):
    """Classeur et fichiers joints (par nom, sans casse), depuis un .xlsx ou un .zip."""
    name = uploaded.name or ""
    if uploaded.size > MAX_ARCHIVE:
        report.add("Le fichier dépasse 250 Mo.")
        report.raise_if_any()
    payload = uploaded.read()
    files = {}
    if name.lower().endswith(".zip"):
        try:
            archive = zipfile.ZipFile(io.BytesIO(payload))
        except zipfile.BadZipFile:
            report.add("Le .zip est illisible.")
            report.raise_if_any()
        entries = [i for i in archive.infolist() if not i.is_dir() and "__MACOSX" not in i.filename]
        if len(entries) > 300 or sum(i.file_size for i in entries) > 2 * MAX_ARCHIVE:
            report.add("Le .zip contient trop de fichiers ou des fichiers trop lourds.")
            report.raise_if_any()
        sheets = [i for i in entries if i.filename.lower().endswith(".xlsx") and not PurePosixPath(i.filename).name.startswith("~$")]
        if len(sheets) != 1:
            report.add("Le .zip doit contenir exactement un tableur .xlsx.")
            report.raise_if_any()
        workbook_bytes = archive.read(sheets[0])
        for entry in entries:
            if entry is not sheets[0]:
                files[PurePosixPath(entry.filename).name.casefold()] = (entry, archive)
    elif name.lower().endswith(".xlsx"):
        workbook_bytes = payload
    else:
        report.add("Déposez un tableur .xlsx, ou un .zip contenant le tableur et ses médias.")
        report.raise_if_any()
    try:
        book = load_workbook(io.BytesIO(workbook_bytes), data_only=True, read_only=True)
    except Exception:
        report.add("Le tableur est illisible. Enregistrez-le au format Excel (.xlsx).")
        report.raise_if_any()
    return book, files


def read_series(book, skill, report):
    if "Série" not in book.sheetnames or "Contenu" not in book.sheetnames:
        report.add("Le tableur doit contenir les onglets « Série » et « Contenu » du modèle Fluentia.")
        report.raise_if_any()
    values = {}
    for row in book["Série"].iter_rows(min_row=2, values_only=True):
        if row and text(row[0]):
            values[text(row[0])] = text(row[1] if len(row) > 1 else "")
    declared = LABELS.get(values.get("Épreuve", "").casefold())
    if declared is None:
        report.add("Indiquez l’épreuve dans l’onglet « Série » (Lecture, Écoute, Rédaction ou Oral).", None, "Épreuve")
    elif declared != skill:
        report.add(f"Ce fichier est un modèle {SKILLS[declared]}, pas {SKILLS[skill]}. Choisissez la bonne épreuve ou le bon fichier.", None, "Épreuve")
    report.raise_if_any()
    for label, limit in (("Titre", 180), ("Provenance et droits", 10000)):
        if not values.get(label):
            report.add(f"« {label} » est obligatoire dans l’onglet « Série ».", None, label)
        elif len(values[label]) > limit:
            report.add(f"« {label} » dépasse {limit} caractères.", None, label)
    mode = ""
    if skill in QCM:
        mode = {"simulation": "simulation", "entraînement": "training", "entrainement": "training"}.get(values.get("Mode", "").casefold(), "")
        if not mode:
            report.add("Mode : choisissez Simulation ou Entraînement.", None, "Mode")
    return {
        "title": values.get("Titre", "")[:180], "mode": mode, "description": values.get("Description", ""),
        "provenance": values.get("Provenance et droits", ""), "credit": values.get("Auteur ou autrice", "")[:250],
    }


def read_rows(book, skill, report):
    sheet = book["Contenu"]
    rows = list(sheet.iter_rows(values_only=True))
    header = [text(v) for v in (rows[0] if rows else [])]
    while header and not header[-1]:
        header.pop()
    expected = COLUMNS[skill]
    if header != expected:
        other = next((s for s, cols in COLUMNS.items() if cols == header), None)
        if other:
            report.add(f"Ces colonnes sont celles du modèle {SKILLS[other]}, pas {SKILLS[skill]}.", 1)
        else:
            missing = [c for c in expected if c not in header]
            report.add("Les titres de colonnes ne correspondent pas au modèle " + SKILLS[skill]
                       + (f" (manquantes : {', '.join(missing)})" if missing else "") + ". Repartez du modèle à télécharger.", 1)
        report.raise_if_any()
    body = []
    for line, row in enumerate(rows[1:], 2):
        cells = {name: (row[i] if i < len(row) else None) for i, name in enumerate(expected)}
        if any(text(v) for v in cells.values()):
            body.append((line, cells))
    if not body:
        report.add("L’onglet « Contenu » ne contient aucune ligne.")
    elif len(body) > MAX_ROWS:
        report.add(f"{MAX_ROWS} lignes maximum par import.")
    report.raise_if_any()
    return body


def attachment(files, name, line, column, report):
    found = files.get(PurePosixPath(name.replace("\\", "/")).name.casefold())
    if not found:
        report.add(f"Fichier « {name} » introuvable dans le .zip.", line, column)
        return None
    entry, archive = found
    if entry.file_size > 25_000_000:
        report.add(f"« {name} » est trop lourd.", line, column)
        return None
    return archive.read(entry)


def check_qcm(skill, meta, body, files, report):
    questions = []
    for expected_number, (line, cells) in enumerate(body, 1):
        def need(column, limit=10000):
            value = text(cells[column])
            if not value:
                report.add(f"« {column} » est vide.", line, column)
            elif len(value) > limit:
                report.add(f"« {column} » dépasse {limit} caractères.", line, column)
            return value

        if text(cells["N°"]) != str(expected_number):
            report.add(f"N° attendu : {expected_number} (les questions doivent se suivre).", line, "N°")
        level = text(cells["Niveau"]).upper()
        if level not in LEVELS:
            report.add("Niveau : A1, A2, B1, B2, C1 ou C2.", line, "Niveau")
        statement = need("Question")
        choices = [need(f"Choix {letter}", 500) for letter in LETTERS]
        if all(choices) and len({c.casefold() for c in choices}) != 4:
            report.add("Les quatre choix doivent être différents.", line, "Choix A")
        answer = text(cells["Bonne réponse"]).upper()
        if answer not in LETTERS or len(answer) != 1:
            report.add("Bonne réponse : A, B, C ou D.", line, "Bonne réponse")
        hint = need("Indice")
        passage = text(cells.get("Texte"))
        if len(passage) > 40000:
            report.add("« Texte » dépasse 40 000 caractères.", line, "Texte")
        image, image_text = None, text(cells["Image"])
        if image_text:
            payload = attachment(files, image_text, line, "Image", report)
            if payload is not None:
                try:
                    check_image(payload)
                    image = (image_text, payload)
                except ValidationError as exc:
                    report.add(f"« {image_text} » : {exc.detail[0] if isinstance(exc.detail, list) else exc.detail}", line, "Image")
            if not text(cells["Texte de l’image"]):
                report.add("« Texte de l’image » est obligatoire quand une image est fournie.", line, "Texte de l’image")
        audio = None
        if skill == "reading" and not passage and not image_text:
            report.add("Il faut un Texte, une Image, ou les deux.", line, "Texte")
        if skill == "listening":
            name = need("Audio", 255)
            transcript = need("Transcription", 40000)
            if name:
                payload = attachment(files, name, line, "Audio", report)
                if payload is not None:
                    try:
                        check_audio(payload)
                        audio = (name, payload, transcript)
                    except ValidationError:
                        report.add(f"« {name} » : un WAV PCM intègre de 15 minutes maximum est requis.", line, "Audio")
        questions.append({
            "line": line, "level": level, "statement": statement, "choices": choices, "answer": answer,
            "hint": hint, "passage": passage, "image": image, "image_text": text(cells["Texte de l’image"]), "audio": audio,
        })
    if meta["mode"] == "simulation":
        if len(questions) != 39:
            report.add(f"Une simulation compte 39 questions ; ce fichier en contient {len(questions)}.")
        ranks = [LEVELS.index(q["level"]) for q in questions if q["level"] in LEVELS]
        if ranks != sorted(ranks):
            report.add("En simulation, les niveaux doivent aller du plus facile au plus difficile (A1 → C2).", None, "Niveau")
    return questions


def check_tasks(skill, body, report):
    subjects = []
    for line, cells in body:
        number_value = text(cells["Tâche"])
        if number_value not in ("1", "2", "3"):
            report.add("Tâche : 1, 2 ou 3.", line, "Tâche")
        for column in ("Titre du sujet", "Consigne"):
            if not text(cells[column]):
                report.add(f"« {column} » est vide.", line, column)
        subject = {"line": line, "task": int(number_value) if number_value in ("1", "2", "3") else 0,
                   "title": text(cells["Titre du sujet"])[:180], "prompt": text(cells["Consigne"])}
        if skill == "writing":
            low, high, minutes = (number(cells[c]) for c in ("Mots minimum", "Mots maximum", "Durée (minutes)"))
            if not low or not high or low >= high:
                report.add("Mots minimum et maximum : deux nombres, le minimum plus petit que le maximum.", line, "Mots minimum")
            if not minutes or minutes <= 0:
                report.add("Durée (minutes) : un nombre positif.", line, "Durée (minutes)")
            if not text(cells["Type de texte"]):
                report.add("« Type de texte » est vide.", line, "Type de texte")
            documents = [text(cells[c]) for c in ("Document 1", "Document 2") if text(cells[c])]
            if subject["task"] == 3 and len(documents) != 2:
                report.add("La tâche 3 demande deux documents à comparer.", line, "Document 1")
            subject.update(genre=text(cells["Type de texte"])[:180], min_words=int(low or 0), max_words=int(high or 0),
                           minutes=int(minutes or 0), documents=documents)
        else:
            prep, speak = number(cells["Préparation (minutes)"]), number(cells["Temps de parole (minutes)"])
            if prep is None or prep < 0:
                report.add("Préparation (minutes) : 0 ou plus.", line, "Préparation (minutes)")
            if not speak or speak <= 0:
                report.add("Temps de parole (minutes) : un nombre positif.", line, "Temps de parole (minutes)")
            raw = text(cells["Pistes"])
            cues = [c.strip() for c in raw.replace(";", "\n").splitlines() if c.strip()]
            subject.update(prep_seconds=round((prep or 0) * 60), speak_seconds=round((speak or 0) * 60), cues=cues)
        subjects.append(subject)
    for number_value in (1, 2, 3):
        group = [s for s in subjects if s["task"] == number_value]
        if not group:
            report.add(f"Aucun sujet pour la tâche {number_value}.")
            continue
        keys = ("genre", "min_words", "max_words", "minutes") if skill == "writing" else ("prep_seconds", "speak_seconds")
        if any(tuple(s[k] for k in keys) != tuple(group[0][k] for k in keys) for s in group):
            report.add(f"Tâche {number_value} : le format (type, mots, durée) doit être identique pour tous ses sujets.", group[0]["line"])
    return subjects


def import_series(actor, skill, uploaded):
    """Vérifie tout le fichier, puis crée la série en brouillon d’un seul tenant."""
    require(actor, "add_bankitem")
    if skill not in SKILLS:
        raise ValidationError({"skill": "Épreuve inconnue."})
    report = Report()
    book, files = open_upload(uploaded, report)
    meta = read_series(book, skill, report)
    body = read_rows(book, skill, report)
    rows = check_qcm(skill, meta, body, files, report) if skill in QCM else check_tasks(skill, body, report)
    report.raise_if_any()
    written = []
    try:
        return save_series(actor, skill, meta, rows, uploaded, written)
    except Exception:
        # La base est annulée par la transaction ; les fichiers déjà écrits sont retirés ici.
        for path in written:
            path.unlink(missing_ok=True)
        raise


@transaction.atomic
def save_series(actor, skill, meta, rows, uploaded, written):
    series = SeriesImport(
        skill=skill, title=meta["title"], mode=meta["mode"], description=meta["description"],
        duration_seconds=DURATIONS.get(skill, 0) * 60, file_name=(uploaded.name or "")[:255], created_by=actor,
    )
    if skill in QCM:
        series.item_ids = [str(create_question(actor, skill, meta, row, written).pk) for row in rows]
    else:
        task_set = TaskSet.objects.create(skill=skill, title=meta["title"], editor=actor, provenance=meta["provenance"])
        positions = {}
        for row in rows:
            positions[row["task"]] = positions.get(row["task"], 0) + 1
            TaskSubject.objects.create(
                task_set=task_set, task_number=row["task"], position=positions[row["task"]], title=row["title"],
                prompt=row["prompt"], genre=row.get("genre", ""), min_words=row.get("min_words"),
                max_words=row.get("max_words"), minutes=row.get("minutes"),
                documents=[{"label": f"Document {i}", "text": t} for i, t in enumerate(row.get("documents", []), 1)],
                prep_seconds=row.get("prep_seconds"), speak_seconds=row.get("speak_seconds"), cues=row.get("cues", []),
            )
        series.task_set = task_set
    series.save()
    return series


def create_question(actor, skill, meta, row, written):
    media = image = None
    if row["audio"]:
        name, payload, transcript = row["audio"]
        media = store_audio(actor, payload, title=f"{meta['title']} · {name}"[:180], transcript=transcript, provenance=meta["provenance"])
        written.append(media_path(media))
    if row["image"]:
        name, payload = row["image"]
        image = store_image(actor, payload, title=f"{meta['title']} · {name}", provenance=meta["provenance"])
        written.append(image_path(image))
    item = create_item(actor, {
        "skill": skill, "difficulty": row["level"], "statement": row["statement"], "passage": row["passage"],
        "media": media, "image": image, "image_text": row["image_text"] if image else "",
        "choices": row["choices"], "correct_index": LETTERS.index(row["answer"]),
        "explanation": f"Indice dans le texte : {row['hint']}.", "origin": "original",
        "provenance": meta["provenance"], "credit": meta["credit"],
    })
    event(item, actor, "import")
    return item

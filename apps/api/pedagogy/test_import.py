"""Recette de l’import Excel : modèles, contrôles ligne par ligne, images, circuit et barème."""

import io
import math
import struct
import tempfile
import wave
import zipfile
import zlib

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import StudentProfile
from exams.models import Exam
from .importer import COLUMNS, build_template
from .models import BankItem, SeriesImport, TaskSet

LEVELS = ["A1"] * 4 + ["A2"] * 6 + ["B1"] * 9 + ["B2"] * 10 + ["C1"] * 6 + ["C2"] * 4  # 699 points


def png(width=800, height=400):
    """PNG valide d’une couleur unie."""
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + b"\xee\xeb\xff" * width for _ in range(height))
    return (bytes([0x89, 0x50, 0x4E, 0x47, 0x0D, 0x0A, 0x1A, 0x0A]) + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def wav(seconds=0.4):
    buffer = io.BytesIO()
    with wave.open(buffer, "wb") as output:
        output.setnchannels(1)
        output.setsampwidth(2)
        output.setframerate(8000)
        samples = [round(800 * math.sin(2 * math.pi * 440 * i / 8000)) for i in range(round(seconds * 8000))]
        output.writeframes(struct.pack("<" + "h" * len(samples), *samples))
    return buffer.getvalue()


def workbook(skill, rows, series=None):
    book = load_workbook(io.BytesIO(build_template(skill)))
    values = {"Titre": "Série de recette", "Mode": "Simulation", "Provenance et droits": "Textes originaux de recette.", **(series or {})}
    for row in book["Série"].iter_rows(min_row=2):
        if row[0].value in values:
            row[1].value = values[row[0].value]
    for row in rows:
        book["Contenu"].append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def reading_rows(count=39, image_on=None):
    rows = []
    for number in range(1, count + 1):
        image = "lettre.png" if number == image_on else None
        rows.append([number, LEVELS[number - 1], None if image else f"Texte original {number}.", image,
                     "Lettre de la voisine du 3e." if image else None, f"Question {number} ?",
                     f"Réponse A{number}", f"Réponse B{number}", f"Réponse C{number}", f"Réponse D{number}", "ABCD"[number % 4], f"extrait {number}"])
    return rows


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class ImportTests(TestCase):
    def setUp(self):
        cache.clear()
        media = tempfile.TemporaryDirectory()
        self.addCleanup(media.cleanup)
        settings = override_settings(PRIVATE_MEDIA_ROOT=media.name)
        settings.enable()
        self.addCleanup(settings.disable)
        User = get_user_model()
        self.editor = User.objects.create_user(username="editrice", is_staff=True, is_superuser=True)
        self.reviewer = User.objects.create_user(username="relecteur", is_staff=True, is_superuser=True)
        self.student = User.objects.create_user(username="eleve", email="eleve@example.test")
        StudentProfile.objects.create(user=self.student, display_name="Alex", email_verified_at=timezone.now(), verified_email=self.student.email)

    def client_for(self, user):
        client = APIClient(enforce_csrf_checks=True)
        client.force_login(user)
        client.csrf = client.get("/api/v1/auth/csrf/").data["csrf_token"]
        return client

    def post(self, user, path, data=None, multipart=False):
        client = self.client_for(user)
        return client.post("/api/v1/" + path, data or {}, format="multipart" if multipart else "json", HTTP_X_CSRFTOKEN=client.csrf)

    def upload(self, skill, payload, name="serie.xlsx", user=None):
        return self.post(user or self.editor, "bank/series/import/", {"skill": skill, "file": SimpleUploadedFile(name, payload)}, multipart=True)

    def run_workflow(self, series_id):
        self.assertEqual(self.post(self.editor, f"bank/series/{series_id}/request_review/").status_code, 200)
        approved = self.post(self.reviewer, f"bank/series/{series_id}/approve/", {"note": "Relu en entier.", "attest": True})
        self.assertEqual(approved.status_code, 200, approved.data)
        return self.post(self.reviewer, f"bank/series/{series_id}/publish/")

    def test_templates_have_three_sheets_and_the_right_columns(self):
        client = self.client_for(self.editor)
        for skill in COLUMNS:
            response = client.get(f"/api/v1/bank/templates/{skill}/")
            self.assertEqual(response.status_code, 200)
            book = load_workbook(io.BytesIO(response.content))
            self.assertEqual(book.sheetnames, ["Série", "Contenu", "Mode d’emploi"])
            self.assertEqual([c.value for c in book["Contenu"][1]], COLUMNS[skill])

    def test_reading_series_goes_from_excel_to_a_699_point_exam(self):
        response = self.upload("reading", workbook("reading", reading_rows()))
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual((response.data["size"], response.data["status"]), (39, "draft"))
        self.assertEqual(BankItem.objects.filter(state="draft").count(), 39)
        # La personne qui importe ne peut pas approuver sa propre série.
        self.post(self.editor, f"bank/series/{response.data['id']}/request_review/")
        self.assertEqual(self.post(self.editor, f"bank/series/{response.data['id']}/approve/", {"note": "Moi.", "attest": True}).status_code, 400)
        published = self.post(self.reviewer, f"bank/series/{response.data['id']}/approve/", {"note": "Relu.", "attest": True})
        self.assertEqual(published.status_code, 200)
        self.assertEqual(self.post(self.reviewer, f"bank/series/{response.data['id']}/publish/").status_code, 200)
        exam = Exam.objects.get(pk=SeriesImport.objects.get().exam_id)
        self.assertEqual((exam.mode, exam.duration_seconds, exam.is_published), ("simulation", 3600, True))
        self.assertEqual(sum(q.points for q in exam.questions.all()), 699)
        self.assertEqual(exam.questions.get(position=39).points, 33)

    def test_new_version_replaces_the_previous_one_in_the_catalogue(self):
        first = self.upload("reading", workbook("reading", reading_rows(3), {"Mode": "Entraînement"}))
        self.assertEqual(self.run_workflow(first.data["id"]).status_code, 200)
        second = self.upload("reading", workbook("reading", reading_rows(3), {"Mode": "Entraînement"}))
        self.assertEqual(self.run_workflow(second.data["id"]).status_code, 200)
        versions = Exam.objects.filter(slug="serie-de-recette").order_by("version")
        self.assertEqual([(e.version, e.is_published) for e in versions], [(1, False), (2, True)])

    def test_wrong_test_type_is_refused_before_anything_is_saved(self):
        response = self.upload("reading", workbook("listening", []))
        self.assertEqual(response.status_code, 400)
        self.assertIn("modèle Écoute, pas Lecture", str(response.data))
        self.assertEqual(SeriesImport.objects.count(), 0)

    def test_errors_are_reported_line_by_line_and_nothing_is_created(self):
        rows = reading_rows(39)
        rows[4][10] = "E"  # ligne 6 du tableur
        rows[7][7] = None  # ligne 9, choix B vide
        response = self.upload("reading", workbook("reading", rows))
        self.assertEqual(response.status_code, 400)
        found = {(e["ligne"], e["colonne"]) for e in response.data["errors"]}
        self.assertIn((6, "Bonne réponse"), found)
        self.assertIn((9, "Choix B"), found)
        self.assertEqual(BankItem.objects.count(), 0)

    def test_simulation_needs_39_questions_in_increasing_levels(self):
        rows = reading_rows(39)
        rows[0][1], rows[38][1] = "C2", "A1"
        response = self.upload("reading", workbook("reading", rows[:38]))
        messages = " ".join(e["message"] for e in response.data["errors"])
        self.assertIn("39 questions", messages)
        self.assertIn("du plus facile au plus difficile", messages)

    def zipped(self, sheet, files):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("serie/serie.xlsx", sheet)
            for name, data in files.items():
                archive.writestr("serie/" + name, data)
        return buffer.getvalue()

    def test_image_question_is_shown_only_to_the_attempt_owner(self):
        sheet = workbook("reading", reading_rows(2, image_on=2), {"Mode": "Entraînement"})
        response = self.upload("reading", self.zipped(sheet, {"lettre.png": png()}), "serie.zip")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(self.run_workflow(response.data["id"]).status_code, 200)
        exam = Exam.objects.get(is_published=True)
        student = self.client_for(self.student)
        attempt = student.post(f"/api/v1/exams/{exam.pk}/attempts/", HTTP_X_CSRFTOKEN=student.csrf).data["id"]
        question = [q for q in student.get(f"/api/v1/attempts/{attempt}/").data["questions"] if q["position"] == 2][0]
        self.assertEqual((question["image_available"], question["image_text"], question["reading_passage"]), (True, "Lettre de la voisine du 3e.", ""))
        image = student.get(f"/api/v1/attempts/{attempt}/questions/{question['id']}/image/")
        self.assertEqual((image.status_code, image["Content-Type"]), (200, "image/png"))
        other = get_user_model().objects.create_user(username="autre")
        self.assertEqual(self.client_for(other).get(f"/api/v1/attempts/{attempt}/questions/{question['id']}/image/").status_code, 404)

    def test_missing_or_tiny_images_are_refused(self):
        sheet = workbook("reading", reading_rows(2, image_on=1), {"Mode": "Entraînement"})
        missing = self.upload("reading", self.zipped(sheet, {}), "serie.zip")
        self.assertIn("introuvable", str(missing.data["errors"]))
        tiny = self.upload("reading", self.zipped(sheet, {"lettre.png": png(120, 80)}), "serie.zip")
        self.assertIn("trop petite", str(tiny.data["errors"]))

    def test_listening_rows_need_their_audio_and_transcription(self):
        row = [1, "A1", "q1.wav", "Bonjour, le train part à 8 h.", None, None, "À quelle heure ?", "7 h", "8 h", "9 h", "10 h", "B", "à 8 h"]
        sheet = workbook("listening", [row], {"Mode": "Entraînement"})
        ok = self.upload("listening", self.zipped(sheet, {"q1.wav": wav()}), "ecoute.zip")
        self.assertEqual(ok.status_code, 201, ok.data)
        item = BankItem.objects.get()
        self.assertEqual((item.skill, item.media.transcript), ("listening", "Bonjour, le train part à 8 h."))
        row[3] = None
        refused = self.upload("listening", self.zipped(workbook("listening", [row], {"Mode": "Entraînement"}), {"q1.wav": wav()}), "ecoute.zip")
        self.assertIn((2, "Transcription"), {(e["ligne"], e["colonne"]) for e in refused.data["errors"]})

    def test_writing_subjects_replace_the_active_set_after_review(self):
        rows = [
            [1, "Fête de rue", "Invitez vos voisins.", "Courriel", 60, 120, 10, None, None],
            [2, "Souvenir", "Racontez un souvenir.", "Récit", 120, 150, 20, None, None],
            [3, "Vélo en ville", "Comparez les deux avis.", "Argumentation", 120, 180, 30, "Pour le vélo.", "Contre le vélo."],
        ]
        response = self.upload("writing", workbook("writing", rows, {"Titre": "Sujets d’automne"}))
        self.assertEqual(response.status_code, 201, response.data)
        before = self.client_for(self.student).get("/api/v1/tasks/writing/").data
        self.assertEqual(before["title"], "Sujets originaux Fluentia")
        self.assertEqual(self.run_workflow(response.data["id"]).status_code, 200)
        after = APIClient().get("/api/v1/tasks/writing/").data
        self.assertEqual(after["title"], "Sujets d’automne")
        self.assertEqual([t["subjects"][0]["title"] for t in after["tasks"]], ["Fête de rue", "Souvenir", "Vélo en ville"])
        self.assertEqual(after["tasks"][2]["subjects"][0]["documents"][1]["text"], "Contre le vélo.")
        self.assertEqual(TaskSet.objects.filter(skill="writing", is_active=True).count(), 1)

    def test_speaking_needs_all_three_tasks(self):
        rows = [[1, "Se présenter", "Parlez de vous.", 0, 2, "Parcours ; loisirs"]]
        response = self.upload("speaking", workbook("speaking", rows))
        messages = " ".join(e["message"] for e in response.data["errors"])
        self.assertIn("Aucun sujet pour la tâche 2", messages)
        self.assertIn("Aucun sujet pour la tâche 3", messages)

    def test_students_cannot_import(self):
        self.assertEqual(self.upload("reading", workbook("reading", reading_rows(1)), user=self.student).status_code, 403)
